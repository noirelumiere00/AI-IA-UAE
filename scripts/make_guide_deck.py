"""AiLa 使い方ガイド（管理職向け）スライド生成。python-pptx・16:9・スピーカーノート付き。
出力: ~/Downloads/AiLa_使い方ガイド.pptx
"""
from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x12, 0x8C, 0x7D)
INK = RGBColor(0x22, 0x2B, 0x35)
GREY = RGBColor(0x5A, 0x66, 0x73)
LIGHT = RGBColor(0xEE, 0xF3, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
JP = "Hiragino Kaku Gothic ProN"  # mac標準。Win等では自動フォールバック

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height


def _jp(run, size, color, bold=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = JP
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", JP)


def _box(slide, x, y, w, h, fill=None, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
    return sp


def _text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=6):
    """runs: list of paragraphs; each = list of (text,size,color,bold)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        for (txt, size, color, bold) in para:
            r = p.add_run()
            r.text = txt
            _jp(r, size, color, bold)
    return tb


def _note(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def content_slide(title, bullets, note, kicker=""):
    """bullets: list of (text, level)  level0=主, level1=従(インデント)。"""
    s = prs.slides.add_slide(BLANK)
    _box(s, 0, 0, W, H, fill=WHITE)
    _box(s, 0, 0, W, Inches(1.25), fill=NAVY)              # タイトル帯
    _box(s, 0, Inches(1.25), W, Pt(5), fill=TEAL)          # アクセント線
    if kicker:
        _text(s, Inches(0.6), Inches(0.18), Inches(11), Inches(0.3),
              [[(kicker, 12, RGBColor(0xBF, 0xD8, 0xE8), True)]])
    _text(s, Inches(0.6), Inches(0.42), Inches(12.1), Inches(0.8),
          [[(title, 27, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    # 本文
    paras = []
    for (txt, lvl) in bullets:
        if lvl == 0:
            paras.append([("●  ", 16, TEAL, True), (txt, 16, INK, False)])
        elif lvl == 1:
            paras.append([("      – ", 13.5, GREY, False), (txt, 13.5, GREY, False)])
        else:  # 強調注記
            paras.append([(txt, 13.5, TEAL, True)])
    _text(s, Inches(0.85), Inches(1.6), Inches(11.7), Inches(5.4), paras, space=9)
    _note(s, note)
    return s


# ── 1. タイトル ───────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
_box(s, 0, 0, W, H, fill=NAVY)
_box(s, 0, Inches(4.55), W, Pt(6), fill=TEAL)
_text(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.6),
      [[("AiLa ｜ 朝のメールサマリー", 44, WHITE, True)]])
_text(s, Inches(0.95), Inches(3.7), Inches(11.5), Inches(0.8),
      [[("使い方ガイド（管理職向け）", 22, RGBColor(0xCF, 0xE2, 0xEC), False)]])
_text(s, Inches(0.95), Inches(6.5), Inches(11.5), Inches(0.5),
      [[("ベクトル株式会社 ・ 2026年6月", 13, RGBColor(0x9F, 0xB6, 0xC6), False)]])
_note(s, "AiLaは『毎朝Slackに今日やることが届くメール秘書』。今日は5分で使い方を説明します。"
         "やることはたった1回の/connectだけ、という所をゴールに。")

# ── 2. 価値 ───────────────────────────────────────────────────────────────────
content_slide(
    "毎朝、Slackに『今日まず何をやるか』が届く",
    [("10秒で把握 ― 未返信・要対応・今日の予定を1通に集約", 0),
     ("失注を防ぐ ― 重要メールの未返信を検知して『今日中に』リマインド", 0),
     ("下書きまで自動 ― 返信文をAIが作成。あなたは確認して送るだけ", 0),
     ("安全第一 ― 勝手に送信しない／見えるのは自分のメールだけ", 0)],
    "メールに埋もれて重要な返信が漏れる・朝の状況把握に時間がかかる、を解決するツール。"
    "特に未返信リマインドは失注防止に直結。", kicker="なぜ AiLa？")

# ── 3. はじめ方 ───────────────────────────────────────────────────────────────
content_slide(
    "はじめ方 ― 設定は1回だけ（所要1分）",
    [("① Slackで AiLa を開き、メッセージ欄に  /connect  と入力", 0),
     ("② 表示されたURLで自分のGoogle（Gmail・カレンダー）を許可", 0),
     ("③ 「✅ 連携が完了しました」と出れば完了 → 翌朝から自動で届く", 0),
     ("見えるのは自分のメールだけ。AiLaが勝手に送信することはありません", 2)],
    "実演できるなら/connectを画面で見せる。『一度きり・1分・自分のメールだけ』を強調。"
    "ここさえ越えれば翌朝から価値が出る。", kicker="STEP")

# ── 4. 読み方 ─────────────────────────────────────────────────────────────────
content_slide(
    "メールサマリーの読み方（上から優先度順）",
    [("🔔 未返信 ― 今日中に … 重要なのに返信が滞留（最優先で対応）", 0),
     ("✉️ 要対応メール … 今日来た『あなた宛』の要返信", 0),
     ("色＝🔴顧客緊急 / 🟠プレス / 🟡顧客 / 🟢社内", 1),
     ("📭 一般 … ニュースレター等は件数だけ（必要な時だけGmail）", 0),
     ("📅 今日の予定 … 時刻・会議室・参加リンク付き", 0),
     ("アイコン：📝下書きあり ⏰締切 ⚠️要確認 ／ 件名タップでGmailが開く", 1)],
    "上から見れば優先度順。色で重要度が一目。件名はGmailへのリンク。"
    "『一般』は件数だけに畳んでノイズを排除している点を説明。", kicker="読み方")

# ── 5. リマインド操作 ─────────────────────────────────────────────────────────
content_slide(
    "未返信リマインドの操作ボタン",
    [("✏️ 対応する … AIが返信下書きを作成（Gmailにも保存）", 0),
     ("✅ 対応済み … リマインド解除（緑✅が付く・↩取り消し可）", 0),
     ("⏰ 後で … 3日後に再通知", 0),
     ("🔕 もう通知しない … 今後このスレッドは通知しない", 0),
     ("あなたが実際に返信すると、リマインドは自動で消えます", 2)],
    "押し間違えても↩取り消せる。対応済みには緑✅スタンプが付くので状態が見える。"
    "返信すれば自動で消える=二重管理不要。", kicker="操作")

# ── 6. 対応する流れ ───────────────────────────────────────────────────────────
content_slide(
    "「対応する」の流れ ― 誤送信はゼロ",
    [("① [✏️ 対応する] → AIが返信下書きを作成", 0),
     ("② スレッドに展開 → [📝 編集] で手直し", 0),
     ("③ [📤 送信] → 確認(1/2) → 確認(2/2) → 送信", 0),
     ("送信は必ず2段階確認。AiLaが勝手に送ることは一切ありません", 2)],
    "ここが安心の核。AIは下書きまで。送信は人が2段階で確認したときだけ。"
    "『AIが暴走して送る』心配は構造的にゼロ、と言い切る。", kicker="返信の流れ")

# ── 7. 安全 ───────────────────────────────────────────────────────────────────
content_slide(
    "安全・プライバシー",
    [("AiLaは下書きを作るだけ。送信は本人が2段確認したときだけ", 0),
     ("見えるのは自分のメール／カレンダーのみ（他人のものは見えない）", 0),
     ("認証情報・カード番号などの秘密は自動でマスキング", 0),
     ("カレンダーの件名を伏せる設定も可能（希望者は管理者へ）", 0)],
    "全社展開で必ず聞かれる安全面を先回り。per-user認可で他人のメールは見えない。"
    "機密の自動マスキングも実装済み。", kicker="安全")

# ── 8. まとめ ─────────────────────────────────────────────────────────────────
content_slide(
    "まとめ ― まずは /connect から",
    [("やることは『/connect』ただ1回だけ", 0),
     ("毎朝、未返信ゼロ・予定の見落としゼロへ", 0),
     ("困ったら：届かない→/connect確認 ／ ボタン無反応→数秒待って再操作", 0),
     ("お問い合わせ：小俣（AiLa 管理者）", 0)],
    "クロージング：『今日この場で/connectしてみましょう』と行動を促す。"
    "1〜2週間使って感想をください、で締める。", kicker="まとめ")

out = os.path.expanduser("~/Downloads/AiLa_使い方ガイド.pptx")
prs.save(out)
print("saved:", out, "/ slides:", len(prs.slides._sldIdLst))
